import io
import re
import unittest
from unittest.mock import patch

import openpyxl
import routes_tail as R
import helpers_shared as H


class ClassExportEnglishTests(unittest.TestCase):
    def setUp(self):
        self.item = dict(category='COC', no=1, issued_date='2026-01-01',
                         description='Repair pipe', due_date='2026-10-01',
                         remark='한글 요약은 제외', action_taken='임시수리 완료', importance='Urgent')
        self.snap = dict(id=1, vessel_id=1, vessel_name_raw='SHIP', class_society='BV', report_date='2026-01-02')

    def query(self, sql, params=(), one=False):
        if 'date(' in sql:
            return {'d': '2026-10-04'}
        if 'class_status_items' in sql:
            return [self.item]
        if 'class_status WHERE id=' in sql:
            return self.snap
        if 'class_status WHERE vessel_id' in sql:
            return [self.snap]
        if 'SELECT name FROM vessels' in sql:
            return {'name': 'SHIP'}
        if 'SELECT id, name FROM vessels' in sql:
            return [{'id': 1, 'name': 'SHIP'}]
        raise AssertionError(sql)

    def export(self, fn, path='/'):
        with R.app.test_request_context(path), patch.object(R, 'query', self.query):
            response = fn.__wrapped__()
            if isinstance(response, tuple):
                return response
            response.direct_passthrough = False
            return openpyxl.load_workbook(io.BytesIO(response.get_data())).active

    def assert_english(self, ws):
        self.assertFalse(any(re.search(r'[가-힣]', str(c.value or '')) for row in ws for c in row))
        self.assertNotIn('한글 요약', [c.value for c in ws[4]])

    def test_single_and_all_translate_actions_but_not_summary_or_english_description(self):
        for fn in (R.api_class_status_export, R.api_class_status_export_all):
            with self.subTest(fn=fn.__name__), patch.object(H, '_translate_texts_en', return_value=['Temporary repair completed']) as tr:
                if fn is R.api_class_status_export:
                    # functools wrapper still needs the route argument.
                    with R.app.test_request_context('/'), patch.object(R, 'query', self.query):
                        response = fn.__wrapped__(1)
                        response.direct_passthrough = False
                        ws = openpyxl.load_workbook(io.BytesIO(response.get_data())).active
                else:
                    ws = self.export(fn)
                self.assert_english(ws)
                tr.assert_called_once_with(['임시수리 완료'])
                headers = [c.value for c in ws[4]]
                self.assertEqual('Repair pipe', ws.cell(5, headers.index('Description') + 1).value)
                self.assertEqual('Temporary repair completed', ws.cell(5, headers.index('Action Taken') + 1).value)
                self.assertEqual(len(headers), ws.max_column)
                self.assertEqual('임시수리 완료', self.item['action_taken'])

    def test_translation_fallback_blocks_download(self):
        with R.app.test_request_context('/'), patch.object(R, 'query', self.query), patch.object(H, '_translate_texts_en', return_value=['임시수리 완료']):
            response, status = R.api_class_status_export.__wrapped__(1)
            self.assertEqual(503, status)
            self.assertIn('English translation unavailable', response.get_json()['error'])

    def test_manager_export_translates_mixed_description_and_keeps_reply_blank(self):
        self.item['description'] = 'Pipe 수리 필요'
        vessel = dict(name='SHIP', class_society='BV', manager='M1', items=[self.item])
        with patch.object(R, '_class_export_vessels', return_value=[vessel]), patch.object(H, '_translate_texts_en', return_value=['Pipe repair required']) as tr:
            ws = self.export(R.api_class_status_export_by_manager, '/?manager=M1')
        self.assert_english(ws)
        self.assertEqual('Pipe repair required', ws.cell(5, 5).value)
        self.assertIsNone(ws.cell(5, 7).value)
        tr.assert_called_once_with(['Pipe 수리 필요'])

    def test_empty_action_and_no_items_need_no_translation(self):
        self.item['action_taken'] = ''
        with patch.object(H, '_translate_texts_en', side_effect=AssertionError('No translation needed')):
            ws = self.export(R.api_class_status_export_all)
        self.assert_english(ws)
        self.assertIsNone(ws.cell(5, 7).value)
        original = self.query
        def no_items(sql, params=(), one=False):
            return [] if 'class_status_items' in sql else original(sql, params, one)
        with R.app.test_request_context('/'), patch.object(R, 'query', no_items):
            response = R.api_class_status_export_all.__wrapped__()
            response.direct_passthrough = False
            ws = openpyxl.load_workbook(io.BytesIO(response.get_data())).active
        self.assert_english(ws)
        self.assertEqual('No open items', ws.cell(5, 5).value)
        self.assertEqual(8, ws.max_column)


if __name__ == '__main__':
    unittest.main()
