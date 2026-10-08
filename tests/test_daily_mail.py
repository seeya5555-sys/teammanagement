"""Daily 업무관리 주간 업데이트 메일 — 서버 불변식 + 맥 러너 회신엑셀 판독 end-to-end."""
import importlib.util
import io
import os
import tempfile
import unittest
from pathlib import Path

import app as appmod

appmod.app.config['CSRF_PROTECT'] = False

RUNNER = Path.home() / '.openclaw/workspace/automation/daily-mail/daily_mail_runner.py'


class DailyMailTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_db = appmod.DATABASE
        self.old_cfg = appmod.app.config['DATABASE']
        db = os.path.join(self.tmp.name, 'test.db')
        appmod.DATABASE = db
        appmod.app.config['DATABASE'] = db
        with appmod.app.app_context():
            appmod.init_db(False)
            from helpers_shared import _get_api_key
            self.key = _get_api_key(create=True)
            from app_core import execute, query
            sup = query('SELECT id FROM supervisors LIMIT 1', one=True)
            sid = sup['id'] if sup else execute("INSERT INTO supervisors(name) VALUES('T')")
            self.vid = execute("INSERT INTO vessels(name, vsl_cd) VALUES('TEST STAR','TSTR')")
            mk = lambda t, p='Normal', s='Open': execute(
                'INSERT INTO issues(supervisor_id, vessel_id, issue_date, item_topic, priority, status) '
                "VALUES(?,?,'2026-09-01',?,?,?)", (sid, self.vid, t, p, s))
            self.i1, self.i2, self.i3 = mk('ME overhaul'), mk('Class COC hull', 'COC & Flag'), mk('Radar', s='InProgress')
            self.closed = mk('Old item', s='Closed')
        self.c = appmod.app.test_client()
        with self.c.session_transaction() as s:
            s['user_id'] = 1
            s['username'] = 'admin'
            s['role'] = 'admin'
        self.h = {'X-API-Key': self.key}

    def tearDown(self):
        appmod.DATABASE = self.old_db
        appmod.app.config['DATABASE'] = self.old_cfg
        self.tmp.cleanup()


    def test_one_substantive_item_means_week_replied(self):
        self._enable();rid=self._claim_sent()
        self.c.post(f'/api/ext/daily-mail/runs/{rid}/issues/{self.i1}/progress', headers=self.h,
                    json={'message_id':'weekly-one', 'progress':'Repair planned tomorrow', 'date':'2026-10-08'})
        self.c.post(f'/api/ext/daily-mail/runs/{rid}/events',headers=self.h,
                    json={'kind':'reply','message_id':'weekly-one'})
        with appmod.app.app_context():
            from daily_mail_service import get_run
            self.assertEqual('replied',get_run(rid)['reply_status'])
            from app_core import query
            self.assertEqual('Open',query('SELECT status FROM issues WHERE id=?',(self.i2,),one=True)['status'])
        next_run=self._claim_sent(week='2026W41')
        with appmod.app.app_context():
            from daily_mail_service import get_run
            self.assertEqual('pending',get_run(next_run)['reply_status'])

    def test_receipt_without_update_stays_pending(self):
        self._enable();rid=self._claim_sent()
        self.c.post(f'/api/ext/daily-mail/runs/{rid}/events',headers=self.h,
                    json={'kind':'reply','message_id':'legacy-receipt','evidence':'Received with thanks'})
        with appmod.app.app_context():
            from daily_mail_service import get_run
            self.assertEqual('pending',get_run(rid)['reply_status'])

    def test_new_row_progress_reply_evidence_counts(self):
        self._enable();rid=self._claim_sent()
        self.c.post(f'/api/ext/daily-mail/runs/{rid}/events',headers=self.h,
                    json={'kind':'reply','message_id':'new-row','payload':{'new_progress':True}})
        with appmod.app.app_context():
            from daily_mail_service import get_run
            self.assertEqual('replied',get_run(rid)['reply_status'])

    def test_raw_source_key_dedups_different_message_and_translation(self):
        self._enable();rid=self._claim_sent()
        url=f'/api/ext/daily-mail/runs/{rid}/issues/{self.i1}/progress'
        body=dict(message_id='outlook:1',progress='수리 진행 중',date='2026-10-06',source_fingerprint='a'*64)
        self.assertFalse(self.c.post(url,headers=self.h,json=body).get_json()['duplicate'])
        second=self.c.post(url,headers=self.h,json=dict(body,message_id='outlook:2',progress='수리 작업 진행 중')).get_json()
        self.assertTrue(second['duplicate'])
        with appmod.app.app_context():
            import json
            from app_core import query
            self.assertEqual(1,len(json.loads(query('SELECT actions FROM issues WHERE id=?',(self.i1,),one=True)['actions'])))
        self.assertFalse(self.c.post(url,headers=self.h,json=dict(body,message_id='new',progress='부품 보급 완료',source_fingerprint='b'*64)).get_json()['duplicate'])
        next_run=self._claim_sent(week='2026W41')
        self.assertFalse(self.c.post(f'/api/ext/daily-mail/runs/{next_run}/issues/{self.i1}/progress',headers=self.h,json=body).get_json()['duplicate'])

    def test_legacy_local_message_still_dedups(self):
        self._enable();rid=self._claim_sent()
        url=f'/api/ext/daily-mail/runs/{rid}/issues/{self.i1}/progress'
        self.c.post(url,headers=self.h,json=dict(message_id='old',progress='Repair started'))
        out=self.c.post(url,headers=self.h,json=dict(message_id='stable-new',legacy_message_id='old',source_fingerprint='a'*64,progress='수리 진행')).get_json()
        self.assertTrue(out['duplicate'])
        same_old=self.c.post(url,headers=self.h,json=dict(message_id='old',source_fingerprint='b'*64,progress='수리 진행')).get_json()
        self.assertTrue(same_old['duplicate'])
        self.assertEqual(400,self.c.post(url,headers=self.h,json=dict(message_id='bad',source_fingerprint='invalid',progress='x')).status_code)

    def test_unique_claim_survives_both_prechecks_missing_row(self):
        from unittest import mock
        import daily_mail_service as service
        self._enable();rid=self._claim_sent()
        url=f'/api/ext/daily-mail/runs/{rid}/issues/{self.i1}/progress'
        body=dict(message_id='first',progress='Repair started',source_fingerprint='c'*64)
        self.c.post(url,headers=self.h,json=body)
        original_query=service.query
        def race_query(sql,*args,**kw):
            if "json_extract(payload,'$.source_fingerprint')" in sql:return None
            return original_query(sql,*args,**kw)
        with mock.patch.object(service,'query',side_effect=race_query):
            second=self.c.post(url,headers=self.h,json=dict(body,message_id='second',progress='수리 시작'))
        self.assertEqual(200,second.status_code);self.assertTrue(second.get_json()['duplicate'])
        # Identical source on another issue must not suppress its independent update.
        other=self.c.post(f'/api/ext/daily-mail/runs/{rid}/issues/{self.i2}/progress',headers=self.h,json=body)
        self.assertFalse(other.get_json()['duplicate'])

    def test_legacy_source_backfill_and_two_real_polls_no_new_action(self):
        from unittest import mock
        from datetime import datetime, timedelta
        import json
        from openpyxl import Workbook
        spec=importlib.util.spec_from_file_location('dedup_daily_runner',RUNNER)
        runner=importlib.util.module_from_spec(spec);spec.loader.exec_module(runner)
        self._enable();rid=self._claim_sent()
        path=Path(self.tmp.name)/'reply.xlsx';workbook=Workbook();sheet=workbook.active
        sheet.append(['Issue ID','Update (reply here)','Status (Open/Closed)'])
        sheet.append([self.i1,'Repair in progress','Open']);workbook.save(path);workbook.close()
        rec=dict(message_id='outlook:changed',subject='RE: [TRMT-DU 2026W40 TSTR] Update',
                 received_at=(datetime.now()+timedelta(seconds=2)).strftime('%Y-%m-%d %H:%M:%S'),
                 sender='master@example.com',body='',attachments=[str(path)])
        self.c.post(f'/api/ext/daily-mail/runs/{rid}/issues/{self.i1}/progress',headers=self.h,
                    json=dict(message_id='old-local',progress='수리 진행 중',date=rec['received_at'][:10]))
        parsed=runner.read_reply_excel(path,{self.i1})[self.i1]['update']
        fp=runner.source_update_fingerprint(rec,parsed,'no')
        with appmod.app.app_context():
            from app_core import query,execute
            event=query("SELECT id,payload FROM daily_mail_events WHERE run_id=? AND kind='update'",(rid,),one=True)
            payload=dict(json.loads(event['payload']),source_fingerprint=fp)
            execute('UPDATE daily_mail_events SET payload=? WHERE id=?',(json.dumps(payload),event['id']))
        def http(method,path,payload=None,**kwargs):
            response=self.c.open(path,method=method,headers=self.h,json=payload)
            return response.status_code,response.get_json()
        from types import SimpleNamespace
        with mock.patch.object(runner,'http',side_effect=http),mock.patch.object(runner,'outlook_scan',return_value=([rec],'')),mock.patch.object(runner,'substantive_attachment_text',return_value=True),mock.patch.object(runner,'to_korean',return_value=['수리 작업 진행 중']) as translation:
            first=runner.cmd_poll(SimpleNamespace(dry=False,no_remind=True))
            second=runner.cmd_poll(SimpleNamespace(dry=False,no_remind=True))
        self.assertTrue(first['ok']);self.assertTrue(first['replies'][0]['log'][0]['dup'])
        self.assertEqual([],second['replies']);self.assertEqual(1,translation.call_count)
        with appmod.app.app_context():
            from app_core import query
            from daily_mail_service import pending_runs
            self.assertEqual(1,len(json.loads(query('SELECT actions FROM issues WHERE id=?',(self.i1,),one=True)['actions'])))
            run=next(x for x in pending_runs() if x['id']==rid)
            self.assertIn(runner.stable_reply_id(rec),run['processed_message_ids'])

    def test_roster_scope_hides_and_blocks_non_assigned_vessels(self):
        with appmod.app.app_context():
            from app_core import execute
            me = execute("INSERT INTO supervisors(name) VALUES('손유석')")
            mine = execute("INSERT INTO vessels(name, vsl_cd) VALUES('MY SHIP','MYSH')")
            execute('INSERT INTO supervisor_vessels(supervisor_id, vessel_id) VALUES(?,?)', (me, mine))
        names = [v['name'] for v in self.c.get('/api/daily-mail/settings').get_json()['vessels']]
        self.assertIn('MY SHIP', names)
        self.assertNotIn('TEST STAR', names)
        r = self.c.put(f'/api/daily-mail/settings/{self.vid}',
                       json={'to_emails': 'master@x.com', 'cc_emails': '', 'enabled': 1})
        self.assertEqual(403, r.status_code)
        cfg = self.c.get('/api/ext/daily-mail/config?week=2026W40', headers=self.h).get_json()
        self.assertNotIn(self.vid, [v['vessel_id'] for v in cfg['vessels']])
        r = self.c.post('/api/ext/daily-mail/runs', headers=self.h,
                        json={'vessel_id': self.vid, 'iso_week': '2026W40', 'issue_ids': [self.i1]})
        self.assertEqual(409, r.status_code)


    def test_dear_name_per_vessel_in_template(self):
        r = self.c.put('/api/daily-mail/template', json={'subject_tpl': '{vessel} open items',
                                                          'body_tpl': 'Dear {dear},\nPlease update {count} items.'})
        self.assertEqual(200, r.status_code, r.get_json())
        r = self.c.put(f'/api/daily-mail/settings/{self.vid}',
                       json={'to_emails': 'master@x.com', 'cc_emails': '', 'enabled': 1, 'dear_name': ' Capt.  Kim '})
        self.assertEqual(200, r.status_code, r.get_json())
        self.assertEqual('Capt. Kim', r.get_json()['dear_name'])
        cfg = self.c.get('/api/ext/daily-mail/config?week=2026W40', headers=self.h).get_json()
        v = next(x for x in cfg['vessels'] if x['vessel_id'] == self.vid)
        self.assertTrue(v['body'].startswith('Dear Capt. Kim,'))
        self.assertEqual('Capt. Kim', v['dear'])
        r = self.c.put(f'/api/daily-mail/settings/{self.vid}',
                       json={'to_emails': 'master@x.com', 'cc_emails': '', 'enabled': 1})   # 필드 누락 = 유지
        self.assertEqual('Capt. Kim', r.get_json()['dear_name'])
        r = self.c.put(f'/api/daily-mail/settings/{self.vid}',
                       json={'to_emails': 'master@x.com', 'cc_emails': '', 'enabled': 1, 'dear_name': ''})
        cfg = self.c.get('/api/ext/daily-mail/config?week=2026W40', headers=self.h).get_json()
        v = next(x for x in cfg['vessels'] if x['vessel_id'] == self.vid)
        self.assertTrue(v['body'].startswith('Dear Sir/Madam,'))
        r = self.c.put(f'/api/daily-mail/settings/{self.vid}',
                       json={'to_emails': 'master@x.com', 'enabled': 1, 'dear_name': '{vessel}'})
        self.assertEqual(400, r.status_code)


    def test_literal_greeting_uses_dear_and_date_format(self):
        from daily_mail_service import render_template, format_mail_date
        self.assertEqual('07th Oct 2026', format_mail_date('2026-10-07'))
        self.assertEqual('01st Nov 2026', format_mail_date('2026-11-01'))
        self.assertEqual('22nd Oct 2026', format_mail_date('2026-10-22'))
        self.assertEqual('13th Oct 2026', format_mail_date('2026-10-13'))
        out = render_template('Dear Sir/Madam,\nreply by {due_date}.', 'V', 1, '2026-10-07', 'Gerasimos')
        self.assertEqual('Dear Gerasimos,\nreply by 07th Oct 2026.', out)
        self.assertTrue(render_template('Dear Sir/Madam,', 'V', 1, '2026-10-07', '').startswith('Dear Sir/Madam,'))


    def test_issue_fingerprint_in_config_and_export(self):
        self._enable()
        cfg = self.c.get('/api/ext/daily-mail/config?week=2026W40', headers=self.h).get_json()
        fp = next(x for x in cfg['vessels'] if x['vessel_id'] == self.vid)['issue_fp']
        r = self.c.get(f'/api/ext/daily-mail/vessels/{self.vid}/export.xlsx?translate=0', headers=self.h)
        self.assertEqual(fp, r.headers.get('X-Issue-Fp'))
        with appmod.app.app_context():
            from app_core import execute
            execute("UPDATE issues SET description='changed' WHERE id=?", (self.i1,))
        cfg = self.c.get('/api/ext/daily-mail/config?week=2026W40', headers=self.h).get_json()
        self.assertNotEqual(fp, next(x for x in cfg['vessels'] if x['vessel_id'] == self.vid)['issue_fp'])


    def test_enabled_toggle_only_changes_enabled(self):
        self._enable()
        r = self.c.post(f'/api/daily-mail/settings/{self.vid}/enabled', json={'enabled': False})
        self.assertEqual(200, r.status_code, r.get_json())
        v = next(x for x in self.c.get('/api/daily-mail/settings').get_json()['vessels'] if x['vessel_id'] == self.vid)
        self.assertEqual((0, 'master@x.com', 'team@x.com'), (v['enabled'], v['to_emails'], v['cc_emails']))
        self.assertEqual(200, self.c.post(f'/api/daily-mail/settings/{self.vid}/enabled', json={'enabled': True}).status_code)
        with appmod.app.app_context():
            from app_core import execute
            vid2 = execute("INSERT INTO vessels(name, vsl_cd) VALUES('NO TO','NOTO')")
        self.assertEqual(400, self.c.post(f'/api/daily-mail/settings/{vid2}/enabled', json={'enabled': True}).status_code)

    def _enable(self):
        r = self.c.put(f'/api/daily-mail/settings/{self.vid}',
                       json={'to_emails': 'master@x.com', 'cc_emails': 'team@x.com; bad', 'enabled': 1})
        self.assertIn(r.status_code, (200, 400))
        if r.status_code == 400:      # 잘못된 주소는 거부되어야 정상 → 정상 주소로 재시도
            r = self.c.put(f'/api/daily-mail/settings/{self.vid}',
                           json={'to_emails': 'master@x.com', 'cc_emails': 'team@x.com', 'enabled': 1})
        self.assertEqual(200, r.status_code, r.get_json())

    def _claim_sent(self, week='2026W40'):
        cfg = self.c.get(f'/api/ext/daily-mail/config?week={week}', headers=self.h).get_json()
        v = next(x for x in cfg['vessels'] if x['vessel_id'] == self.vid)
        self.assertEqual(sorted([self.i1, self.i2, self.i3]), sorted(v['open_issue_ids']))
        self.assertTrue(v['subject'].startswith(f'[TRMT-DU {week} TSTR]'))
        r = self.c.post('/api/ext/daily-mail/runs', headers=self.h, json={
            'vessel_id': self.vid, 'iso_week': week, 'tag': v['tag'], 'issue_ids': v['open_issue_ids']})
        self.assertEqual(201, r.status_code, r.get_json())
        rid = r.get_json()['id']
        self.assertEqual('sent', self.c.post(f'/api/ext/daily-mail/runs/{rid}/state', headers=self.h,
                                             json={'state': 'sent'}).get_json()['state'])
        return rid

    def test_new_issue_from_reply_row_is_created_once(self):
        self._enable()
        rid = self._claim_sent()
        body = {'message_id': 'm-new', 'issue_date': '2026-09-04', 'due_date': '2026-10-15',
                'item_en': 'SAT C printer not working', 'item_topic': 'SAT-C printer 작동 불량',
                'description': '1. SAT-C printer 작동 불량', 'progress': 'Troubleshooting 진행 중', 'priority': 'Bogus'}
        r = self.c.post(f'/api/ext/daily-mail/runs/{rid}/new-issues', headers=self.h, json=body)
        self.assertEqual(200, r.status_code, r.get_json())
        first = r.get_json()
        self.assertFalse(first['duplicate'])
        # 다른 메일·표기 차이(대소문자/구두점)여도 같은 발생일+Item 이면 재생성 안 함
        again = dict(body, message_id='m-other', item_en='SAT-C Printer not working!')
        r2 = self.c.post(f'/api/ext/daily-mail/runs/{rid}/new-issues', headers=self.h, json=again).get_json()
        self.assertEqual({'issue_id': first['issue_id'], 'duplicate': True}, r2)
        with appmod.app.app_context():
            from app_core import query
            row = query('SELECT * FROM issues WHERE id=?', (first['issue_id'],), one=True)
            self.assertEqual((self.vid, 'Open', 'Normal', '2026-09-04', '2026-10-15', 'daily-mail'),
                             (row['vessel_id'], row['status'], row['priority'], row['issue_date'], row['due_date'], row['created_by']))
            self.assertIn('Troubleshooting', row['actions'])
            self.assertEqual(1, query('SELECT COUNT(*) n FROM daily_mail_new_issues', one=True)['n'])
        hi = self.c.post(f'/api/ext/daily-mail/runs/{rid}/new-issues', headers=self.h,
                         json=dict(body, item_en='Main Air Compressor low efficiency', priority=' High ')).get_json()
        with appmod.app.app_context():
            from app_core import query
            self.assertEqual('Urgent', query('SELECT priority FROM issues WHERE id=?', (hi['issue_id'],), one=True)['priority'])
        self.assertEqual(400, self.c.post(f'/api/ext/daily-mail/runs/{rid}/new-issues', headers=self.h,
                                          json=dict(body, issue_date='Sep 4')).status_code)
        self.assertNotEqual(200, self.c.post(f'/api/ext/daily-mail/runs/{rid}/new-issues', json=body).status_code)

    def test_requires_api_key_and_disabled_vessel_not_listed(self):
        self.assertEqual(401, self.c.get('/api/ext/daily-mail/config').status_code)
        cfg = self.c.get('/api/ext/daily-mail/config', headers=self.h).get_json()
        self.assertFalse([v for v in cfg['vessels'] if v['vessel_id'] == self.vid])

    def test_one_run_per_vessel_week_and_scope(self):
        self._enable()
        rid = self._claim_sent()
        dup = self.c.post('/api/ext/daily-mail/runs', headers=self.h, json={
            'vessel_id': self.vid, 'iso_week': '2026W40', 'issue_ids': [self.i1]})
        self.assertEqual(409, dup.status_code)
        foreign = self.c.post('/api/ext/daily-mail/runs', headers=self.h, json={
            'vessel_id': self.vid, 'iso_week': '2026W41', 'issue_ids': [self.closed]})
        self.assertEqual(409, foreign.status_code)
        out = self.c.post(f'/api/ext/daily-mail/runs/{rid}/issues/{self.closed}/progress', headers=self.h,
                          json={'message_id': 'm1', 'progress': 'x'})
        self.assertEqual(409, out.status_code)

    def test_update_close_coc_suggest_and_reopen(self):
        self._enable()
        rid = self._claim_sent()
        p = lambda iid, kind, body: self.c.post(f'/api/ext/daily-mail/runs/{rid}/issues/{iid}/{kind}',
                                                headers=self.h, json=body)
        r = p(self.i1, 'progress', {'message_id': 'm1', 'progress': 'Parts arrived', 'sender': 'master@x.com'})
        self.assertEqual(200, r.status_code)
        self.assertTrue(p(self.i1, 'progress', {'message_id': 'm1', 'progress': 'Parts arrived'}).get_json()['duplicate'])
        coc = p(self.i2, 'close', {'message_id': 'm1', 'evidence': 'done'})
        self.assertEqual(409, coc.status_code)
        self.assertEqual('needs_suggestion', coc.get_json().get('code'))
        self.assertEqual(200, p(self.i2, 'close-suggest', {'message_id': 'm1', 'evidence': 'done'}).status_code)
        self.assertEqual(200, p(self.i3, 'close', {'message_id': 'm1', 'evidence': 'Radar repaired'}).status_code)
        with appmod.app.app_context():
            from app_core import query
            self.assertEqual('Closed', query('SELECT status FROM issues WHERE id=?', (self.i3,), one=True)['status'])
            acts = query('SELECT actions FROM issues WHERE id=?', (self.i1,), one=True)['actions']
            self.assertIn('Parts arrived', acts)
            n_issues = query('SELECT COUNT(*) n FROM issues', one=True)['n']
        st = self.c.get('/api/daily-mail/status?week=2026W40').get_json()
        self.assertEqual(1, len(st['suggestions']))
        close_ev = st['closes'][0]['id']
        self.assertEqual(200, self.c.post(f'/api/daily-mail/closes/{close_ev}/reopen').status_code)
        with appmod.app.app_context():
            from app_core import query
            self.assertEqual('InProgress', query('SELECT status FROM issues WHERE id=?', (self.i3,), one=True)['status'])
            self.assertEqual(n_issues, query('SELECT COUNT(*) n FROM issues', one=True)['n'])   # INSERT 없음

    def test_moved_issue_and_reminder_claim_dedup(self):
        self._enable()
        rid = self._claim_sent()
        with appmod.app.app_context():
            from app_core import execute
            other = execute("INSERT INTO vessels(name, vsl_cd) VALUES('OTHER','OTHR')")
            execute('UPDATE issues SET vessel_id=? WHERE id=?', (other, self.i1))
        r = self.c.post(f'/api/ext/daily-mail/runs/{rid}/issues/{self.i1}/progress', headers=self.h,
                        json={'message_id': 'm1', 'progress': 'x'})
        self.assertEqual(409, r.status_code)
        ev = lambda: self.c.post(f'/api/ext/daily-mail/runs/{rid}/events', headers=self.h,
                                 json={'kind': 'reminder', 'message_id': 'reminder-1'}).get_json()
        self.assertFalse(ev()['duplicate'])
        self.assertTrue(ev()['duplicate'])
        self.assertEqual(409, self.c.post(f'/api/daily-mail/runs/{rid}/release').status_code)

    @unittest.skipUnless(RUNNER.exists(), 'mac runner not present')
    def test_runner_reads_filled_reply_excel(self):
        self._enable()
        r = self.c.get(f'/api/ext/daily-mail/vessels/{self.vid}/export.xlsx?translate=0', headers=self.h)
        self.assertEqual(200, r.status_code)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.data))
        ws = wb.worksheets[0]
        hdr_row, cols = None, None
        for row in ws.iter_rows():
            vals = [c.value for c in row]
            if 'Issue ID' in vals:
                hdr_row, cols = row[0].row, {v: i + 1 for i, v in enumerate(vals) if v}
                break
        self.assertIsNotNone(hdr_row)
        for rr in range(hdr_row + 1, ws.max_row + 1):
            iid = ws.cell(rr, cols['Issue ID']).value
            if iid == self.i1:
                ws.cell(rr, cols['Update (reply here)']).value = 'Overhaul finished 30 Sep'
                ws.cell(rr, cols['Status (Open/Closed)']).value = 'Closed'
        path = os.path.join(self.tmp.name, 'reply.xlsx')
        wb.save(path)
        spec = importlib.util.spec_from_file_location('dmr', RUNNER)
        dmr = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(dmr)
        got = dmr.read_reply_excel(path, {self.i1, self.i2, self.i3})
        self.assertEqual({self.i1}, set(got))
        self.assertTrue(got[self.i1]['closed'])
        self.assertEqual('Overhaul finished 30 Sep', got[self.i1]['update'])
        self.assertEqual('Thanks', dmr.top_segment('Thanks\n\nFrom: TSI\nSent: x\nold'))



class DailyMailNavTests(unittest.TestCase):
    def test_admin_nav_links_to_settings_page(self):
        base = (Path(__file__).resolve().parent.parent / 'templates' / 'base.html').read_text(encoding='utf-8')
        self.assertEqual(base.count("nlink('routes_daily_mail.daily_mail_page', 'Daily')"), 2)
        self.assertIn("'routes_daily_mail.daily_mail_page','routes_vetting_mail.vetting_mail_page','routes_class_mail.class_mail_page'] %}", base)

if __name__ == '__main__':
    unittest.main()
