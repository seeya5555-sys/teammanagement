"""Condition Survey (CONTAINER) mail automation: server contract tests."""
import base64
import contextlib
import datetime as dt
import hashlib
import io
import json
import os
import re
import tempfile
import time
import unittest
import zipfile
from datetime import date, datetime
from types import SimpleNamespace
from unittest.mock import patch

import app as A
import cs_mail_service as S
import class_mail_service as CLASS
import mail_reminder_service as R
from app_core import execute, query


class _Frozen(dt.datetime):
    @classmethod
    def now(cls, tz=None):
        return dt.datetime(2026, 10, 5, 1, 2, 3, tzinfo=tz)


@contextlib.contextmanager
def frozen_xlsx_clock():
    """openpyxl stamps docProps and zip entries with wall time; freeze both for byte equality."""
    import openpyxl.packaging.core as core
    import openpyxl.writer.excel as writer
    fake_dt = SimpleNamespace(datetime=_Frozen, timezone=dt.timezone)
    fake_time = SimpleNamespace(time=lambda: 1791000000.0, localtime=time.localtime)
    with patch.object(core, 'datetime', fake_dt), patch.object(writer, 'datetime', fake_dt), patch.object(zipfile, 'time', fake_time):
        yield


class CsMailTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_db, self.old_cfg = A.DATABASE, A.app.config['DATABASE']
        self.old_csrf = A.app.config.get('CSRF_PROTECT')
        A.DATABASE = os.path.join(self.temp.name, 'test.db')
        A.app.config['DATABASE'] = A.DATABASE
        A.app.config['CSRF_PROTECT'] = False
        with A.app.app_context(), contextlib.redirect_stdout(io.StringIO()):
            A.init_db(False)
            from helpers_shared import _get_api_key
            self.key = _get_api_key(create=True)
            sup = query("SELECT id FROM supervisors WHERE name='손유석'", one=True)
            sid = sup['id'] if sup else execute("INSERT INTO supervisors(name) VALUES('손유석')")
            self.vid = execute("INSERT INTO vessels(name,vsl_cd,vessel_type,active) VALUES('BOX STAR','BXST','CNTR',1)")
            self.tanker = execute("INSERT INTO vessels(name,vsl_cd,vessel_type,active) VALUES('TANK STAR','TKST','VLCC',1)")
            self.idle = execute("INSERT INTO vessels(name,vsl_cd,vessel_type,active) VALUES('IDLE BOX','IDBX','CONTAINER',1)")
            for vid in (self.vid, self.tanker, self.idle):
                execute('INSERT INTO supervisor_vessels(supervisor_id,vessel_id) VALUES(?,?)', (sid, vid))
                execute("INSERT INTO daily_mail_settings(vessel_id,to_emails,cc_emails,dear_name,enabled) VALUES(?,'captain@example.com','team@example.com','Capt. Test',0)", (vid,))
            self.s_old = execute("INSERT INTO cs_surveys(vessel_id,year,quarter,vendor,inspection_date) VALUES(?,2025,4,'IDWAL','2025-11-02')", (self.vid,))
            self.s_new = execute("INSERT INTO cs_surveys(vessel_id,year,quarter,vendor,inspection_date) VALUES(?,2026,2,'AALMAR','2026-05-10')", (self.vid,))
            self.s_done = execute("INSERT INTO cs_surveys(vessel_id,year,quarter) VALUES(?,2026,1)", (self.vid,))
            self.f_old = execute("INSERT INTO cs_findings(survey_id,category,no,item,description,remark,status) VALUES(?,'Observation',1,'Hatch cover','Rubber packing damaged','해치 패킹 손상','Open')", (self.s_old,))
            self.f1 = execute("INSERT INTO cs_findings(survey_id,category,no,item,description,remark,status) VALUES(?,'Defect',1,'Mooring winch','Brake lining worn','브레이크 마모','Open')", (self.s_new,))
            self.f2 = execute("INSERT INTO cs_findings(survey_id,category,no,item,description,remark,status) VALUES(?,'Observation',1,'Deck paint','Rust on main deck','','Closed')", (self.s_new,))
            self.f3 = execute("INSERT INTO cs_findings(survey_id,category,no,item,description,remark,status) VALUES(?,'Observation',2,'Lifeboat','Davit wire due renewal','','Open')", (self.s_new,))
            execute("INSERT INTO cs_findings(survey_id,category,no,item,description,status) VALUES(?,'Defect',1,'Old','Closed only','Closed')", (self.s_done,))
            ts = execute("INSERT INTO cs_surveys(vessel_id,year,quarter) VALUES(?,2026,2)", (self.tanker,))
            execute("INSERT INTO cs_findings(survey_id,category,no,item,description,status) VALUES(?,'Defect',1,'Pump','Leak','Open')", (ts,))
        self.client = A.app.test_client()
        with self.client.session_transaction() as session:
            session.update(user_id=1, role='admin', username='admin')
        self.headers = {'X-API-Key': self.key}
        self.clock = patch.object(S, 'today_kst', return_value=date(2026, 10, 5))   # Monday W41
        self.clock.start()

    def tearDown(self):
        self.clock.stop()
        A.DATABASE, A.app.config['DATABASE'] = self.old_db, self.old_cfg
        A.app.config['CSRF_PROTECT'] = self.old_csrf
        self.temp.cleanup()

    def app_context(self):
        return A.app.app_context()

    def get(self, path):
        return self.client.get('/api/ext/cs-mail' + path, headers=self.headers)

    def post(self, path, body):
        return self.client.post('/api/ext/cs-mail' + path, headers=self.headers, json=body)

    def claim_body(self, prepared):
        return {'vessel_id': self.vid, 'iso_week': prepared['iso_week'], 'fingerprint': prepared['fingerprint'],
                'excel_sha256': prepared['excel_sha256'], 'reply_rows': prepared['reply_rows'],
                'attachments': [{k: f[k] for k in ('survey_id', 'filename', 'sha256')} for f in prepared['files']]}

    def prepare_and_claim(self):
        self.client.post(f'/api/cs-mail/settings/{self.vid}/enabled', json={'enabled': True})
        prepared = self.get(f'/vessels/{self.vid}/prepare').get_json()
        return self.post('/runs', self.claim_body(prepared)), prepared

    def sent_run(self, sent_at='2026-10-05 07:31:00'):
        response, _ = self.prepare_and_claim()
        self.assertEqual(201, response.status_code, response.get_json())
        rid = response.get_json()['id']
        self.assertEqual(200, self.post(f'/runs/{rid}/state', {'state': 'sent'}).status_code)
        with self.app_context():
            execute('UPDATE cs_mail_runs SET sent_at=? WHERE id=?', (sent_at, rid))
        return rid

    # ── targets ──
    def test_container_only_roster_default_off_and_no_open_skipped(self):
        body = self.client.get('/api/cs-mail/status').get_json()
        names = {v['name']: v for v in body['vessels']}
        self.assertNotIn('TANK STAR', names)                       # VLCC excluded
        self.assertEqual({'BOX STAR', 'IDLE BOX'}, set(names))    # CNTR and CONTAINER spellings
        self.assertEqual(0, names['BOX STAR']['enabled'])
        self.assertEqual(3, names['BOX STAR']['count'])
        self.assertEqual(2, names['BOX STAR']['survey_count'])     # all-closed survey skipped
        self.assertEqual(0, names['IDLE BOX']['count'])
        self.assertEqual([], self.get('/config').get_json()['vessels'])
        self.client.post(f'/api/cs-mail/settings/{self.idle}/enabled', json={'enabled': True})
        self.client.post(f'/api/cs-mail/settings/{self.vid}/enabled', json={'enabled': True})
        self.assertEqual(['BOX STAR'], [v['name'] for v in self.get('/config').get_json()['vessels']])
        self.assertEqual(403, self.client.post(f'/api/cs-mail/settings/{self.tanker}/enabled', json={'enabled': True}).status_code)
        self.assertEqual(404, self.get(f'/vessels/{self.idle}/prepare').status_code)
        with self.app_context():
            execute("DELETE FROM supervisors WHERE name='손유석'")
        self.assertEqual([], self.get('/config').get_json()['vessels'])
        self.assertEqual(401, A.app.test_client().get('/api/ext/cs-mail/config').status_code)

    def test_distinct_tag_never_matches_class_scan(self):
        vessel = next(v for v in self.client.get('/api/cs-mail/status').get_json()['vessels'] if v['vessel_id'] == self.vid)
        self.assertTrue(vessel['subject'].startswith(f'[TRMT-COND 2026W41 V{self.vid}] BOX STAR - Condition Survey'))
        self.assertNotIn('[TRMT-CS ', vessel['subject'])
        self.assertFalse(CLASS.TAG_RE.search(vessel['subject']))
        self.assertFalse(S.TAG_RE.search('[TRMT-CS 2026W41 V1] X - Class Status'))
        self.assertIn('RECTIFICATION', vessel['body'])
        self.assertIn('Dear Capt. Test,', vessel['body'])
        self.assertFalse(S.KOREAN.search(vessel['subject'] + vessel['body']))

    # ── attachments ──
    def test_attachments_byte_equal_to_export_button(self):
        self.client.post(f'/api/cs-mail/settings/{self.vid}/enabled', json={'enabled': True})
        with frozen_xlsx_clock():
            prepared = self.get(f'/vessels/{self.vid}/prepare').get_json()
            self.assertEqual([self.s_old, self.s_new], [f['survey_id'] for f in prepared['files']])
            for item in prepared['files']:
                exported = self.client.get(f"/api/cs/surveys/{item['survey_id']}/export")
                self.assertEqual(200, exported.status_code)
                raw = base64.b64decode(item['xlsx_base64'])
                self.assertEqual(exported.data, raw)                          # byte-level identical
                self.assertEqual(hashlib.sha256(raw).hexdigest(), item['sha256'])
                self.assertIn(item['filename'], exported.headers['Content-Disposition'])
        self.assertEqual(['CS_BOX STAR_2025Q4.xlsx', 'CS_BOX STAR_2026Q2.xlsx'], [f['filename'] for f in prepared['files']])
        # Unfrozen: everything except the timestamp part is still identical.
        prepared = self.get(f'/vessels/{self.vid}/prepare').get_json()
        exported = self.client.get(f'/api/cs/surveys/{self.s_new}/export').data
        mine = zipfile.ZipFile(io.BytesIO(base64.b64decode(prepared['files'][1]['xlsx_base64'])))
        theirs = zipfile.ZipFile(io.BytesIO(exported))
        self.assertEqual(mine.namelist(), theirs.namelist())
        for name in mine.namelist():
            if name != 'docProps/core.xml':
                self.assertEqual(mine.read(name), theirs.read(name), name)
        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(exported)).active
        self.assertEqual(['Category', 'No.', 'ITEM', 'DESCRIPTION', 'RECTIFICATION', 'PHOTO', 'STATUS'], [c.value for c in ws[4]])
        self.assertEqual([self.f1, self.f2, self.f3], [r['finding_id'] for r in prepared['reply_rows'] if r['survey_id'] == self.s_new])

    def test_claim_duplicate_stale_and_server_owned_recipients(self):
        response, prepared = self.prepare_and_claim()
        self.assertEqual(201, response.status_code, response.get_json())
        run = response.get_json()
        self.assertEqual('captain@example.com', run['to_emails'])
        self.assertEqual(2, len(run['attachments']))
        self.assertEqual(409, self.post('/runs', {**self.claim_body(prepared), 'to_emails': 'x@example.com'}).status_code)
        with self.app_context():
            execute('DELETE FROM cs_mail_item_state')
            execute('DELETE FROM cs_mail_runs')
            execute("UPDATE cs_findings SET description='Brake lining worn out' WHERE id=?", (self.f1,))
        self.assertEqual(409, self.post('/runs', self.claim_body(prepared)).status_code)
        bad = self.get(f'/vessels/{self.vid}/prepare').get_json()
        tampered = self.claim_body(bad)
        tampered['attachments'][0]['sha256'] = '0' * 64
        self.assertEqual(400, self.post('/runs', tampered).status_code)
        with patch.object(S, 'today_kst', return_value=date(2026, 10, 6)):   # Tuesday: not eligible
            tue = self.get(f'/vessels/{self.vid}/prepare').get_json()
            self.assertEqual(409, self.post('/runs', self.claim_body(tue)).status_code)

    def test_pre_send_check_detects_contact_change(self):
        response, _ = self.prepare_and_claim()
        rid = response.get_json()['id']
        self.assertEqual(200, self.get(f'/runs/{rid}/check').status_code)
        with self.app_context():
            execute("UPDATE daily_mail_settings SET cc_emails='changed@example.com' WHERE vessel_id=?", (self.vid,))
        self.assertEqual(409, self.get(f'/runs/{rid}/check').status_code)

    def test_recipient_override_follows_daily_by_default_and_resets(self):
        url = f'/api/cs-mail/settings/{self.vid}/recipients'
        vessel = lambda: next(v for v in self.client.get('/api/cs-mail/status').get_json()['vessels'] if v['vessel_id'] == self.vid)
        self.assertEqual('captain@example.com', vessel()['to_emails'])
        self.assertFalse(any(vessel()['recipients_custom'].values()))
        r = self.client.put(url, json={'to_emails': 'cs@example.com, b@example.com', 'cc_emails': '', 'dear_name': 'Chief'})
        self.assertEqual(200, r.status_code, r.get_json())
        v = vessel()
        self.assertEqual(('cs@example.com; b@example.com', '', 'Chief'), (v['to_emails'], v['cc_emails'], v['dear_name']))
        self.assertIn('Dear Chief', v['body'])
        with self.app_context():   # Daily 변경은 override 된 선박에 영향 없음
            execute("UPDATE daily_mail_settings SET to_emails='other@example.com' WHERE vessel_id=?", (self.vid,))
        self.assertEqual('cs@example.com; b@example.com', vessel()['to_emails'])
        # 부분 수정: 누락 키는 유지, 명시 null 은 해당 칸만 Daily 로
        self.assertEqual(200, self.client.put(url, json={'dear_name': 'Master'}).status_code)
        self.assertEqual(('cs@example.com; b@example.com', 'Master'), (vessel()['to_emails'], vessel()['dear_name']))
        self.client.put(url, json={'cc_emails': None})
        self.assertFalse(vessel()['recipients_custom']['cc_emails'])
        self.assertTrue(vessel()['recipients_custom']['to_emails'])
        # 발송 claim 은 실효(CS 전용) 수신처 사용
        self.client.post(f'/api/cs-mail/settings/{self.vid}/enabled', json={'enabled': True})
        response, _ = self.prepare_and_claim()
        self.assertEqual('cs@example.com; b@example.com', response.get_json()['to_emails'])
        rid = response.get_json()['id']
        self.client.put(url, json={'to_emails': 'cs@example.com'})
        self.assertEqual(409, self.get(f'/runs/{rid}/check').status_code)
        # 등록된 CS 전용 수신처의 회신은 신뢰
        with self.app_context():
            self.assertTrue(S.sender_trusted({'vessel_id': self.vid, 'to_emails': '', 'cc_emails': ''}, 'CS <cs@example.com>'))
        self.assertEqual(400, self.client.put(url, json={'to_emails': 'not-an-email'}).status_code)
        self.assertEqual(400, self.client.put(url, json={'dear_name': '선장'}).status_code)
        self.assertEqual(403, self.client.put(f'/api/cs-mail/settings/{self.tanker}/recipients', json={'to_emails': 'a@example.com'}).status_code)
        # ON 상태에서 To 를 비우면 OFF 로 내려감
        self.client.post(f'/api/cs-mail/settings/{self.vid}/enabled', json={'enabled': True})
        r = self.client.put(url, json={'to_emails': ''}).get_json()
        self.assertEqual(0, r['enabled'])
        self.assertTrue(r['blocker'])
        self.assertEqual(200, self.client.put(url, json={'reset': True}).status_code)
        v = vessel()
        self.assertEqual('other@example.com', v['to_emails'])
        self.assertFalse(any(v['recipients_custom'].values()))

    def test_quarter_variable_rendered(self):
        self.assertEqual('2026 Q1, Q3 / 2027 Q1', S.quarter_label([{'year': 2026, 'quarter': 1}, {'year': 2026, 'quarter': 3}, {'year': 2027, 'quarter': 1}]))
        with self.app_context():
            v = next(v for v in S.settings() if v['vessel_id'] == self.vid)
        label = S.quarter_label(v['surveys'])
        self.assertTrue(label)
        self.assertIn(f'({label})', v['subject'])
        self.assertIn(f'({label})', v['body'])
        self.assertNotIn('{quarter}', v['subject'] + v['body'])

    # ── replies ──
    def test_reply_remark_dated_dedup_close_only_when_flagged(self):
        rid = self.sent_run()
        body = {'message_id': 'm1', 'sender': 'captain@example.com', 'received_at': '2026-10-06 10:00:00',
                'updates': [{'finding_id': self.f1, 'remark': 'Brake lining renewed and tested satisfactory.', 'close': True},
                            {'finding_id': self.f3, 'remark': 'Wire ordered, to be renewed at Busan.', 'close': False}]}
        out = self.post(f'/runs/{rid}/reply', body).get_json()
        self.assertEqual((2, 1), (out['updated'], out['closed']), out)
        self.assertTrue(self.post(f'/runs/{rid}/reply', body).get_json()['duplicate'])
        with self.app_context():
            f1 = query('SELECT * FROM cs_findings WHERE id=?', (self.f1,), one=True)
            f3 = query('SELECT * FROM cs_findings WHERE id=?', (self.f3,), one=True)
            self.assertEqual('브레이크 마모\n\n[10/6]Brake lining renewed and tested satisfactory.', f1['remark'])
            self.assertEqual('Closed', f1['status'])
            self.assertEqual('[10/6]Wire ordered, to be renewed at Busan.', f3['remark'])
            self.assertEqual('Open', f3['status'])
        # Same text from another sender (copy) → no new line, still counts as replied.
        copy = {**body, 'message_id': 'm2', 'sender': 'team@example.com',
                'updates': [{'finding_id': self.f3, 'remark': 'Wire  ordered, to be renewed at Busan.', 'close': False}]}
        self.assertEqual((0, 1), (lambda o: (o['updated'], o['unchanged']))(self.post(f'/runs/{rid}/reply', copy).get_json()))
        # Closed-at-send or unknown finding cannot be touched; Closed never reopens.
        for fid in (self.f2, 99999):
            self.assertEqual(400, self.post(f'/runs/{rid}/reply', {**body, 'message_id': 'x', 'updates': [{'finding_id': fid, 'remark': 'Done', 'close': True}]}).status_code)
        again = {**body, 'message_id': 'm3', 'updates': [{'finding_id': self.f1, 'remark': 'Re-checked, in progress', 'close': False}]}
        self.post(f'/runs/{rid}/reply', again)
        with self.app_context():
            self.assertEqual('Closed', query('SELECT status FROM cs_findings WHERE id=?', (self.f1,), one=True)['status'])

    def test_concurrent_manual_remark_edit_blocks_without_overwrite(self):
        rid = self.sent_run()
        with self.app_context():
            execute("UPDATE cs_findings SET remark='manual' WHERE id=?", (self.f1,))
        response = self.post(f'/runs/{rid}/reply', {'message_id': 'c1', 'sender': 'captain@example.com',
                                                     'updates': [{'finding_id': self.f1, 'remark': 'Completed', 'close': True}]})
        self.assertEqual(409, response.status_code)
        with self.app_context():
            row = query('SELECT remark,status FROM cs_findings WHERE id=?', (self.f1,), one=True)
            self.assertEqual(('manual', 'Open'), (row['remark'], row['status']))
            self.assertEqual(0, query('SELECT COUNT(*) n FROM cs_mail_replies', one=True)['n'])

    def test_unregistered_sender_is_held_without_any_change(self):
        rid = self.sent_run()
        out = self.post(f'/runs/{rid}/reply', {'message_id': 'u1', 'sender': 'Stranger <stranger@example.org>',
                        'updates': [{'finding_id': self.f1, 'remark': 'Completed.', 'close': True}]}).get_json()
        self.assertEqual(('회신 확인 필요(발신자 미등록)', 0), (out['held'], out['updated']))
        with self.app_context():
            row = query('SELECT remark,status FROM cs_findings WHERE id=?', (self.f1,), one=True)
            self.assertEqual(('브레이크 마모', 'Open'), (row['remark'], row['status']))
            self.assertEqual('회신 확인 필요(발신자 미등록)', query('SELECT note FROM cs_mail_replies', one=True)['note'])
        # Registered CC (display-name form) is trusted.
        ok = self.post(f'/runs/{rid}/reply', {'message_id': 'u2', 'sender': 'Team <TEAM@example.com>',
                       'updates': [{'finding_id': self.f1, 'remark': 'Completed.', 'close': True}]}).get_json()
        self.assertEqual((1, 1), (ok['updated'], ok['closed']))

    def test_manual_status_change_after_send_is_not_overridden(self):
        rid = self.sent_run()
        with self.app_context():
            execute("UPDATE cs_findings SET status='Closed' WHERE id=?", (self.f3,))   # closed by hand
            execute("UPDATE cs_findings SET status='Closed' WHERE id=?", (self.f1,))
            execute("UPDATE cs_findings SET status='Open' WHERE id=?", (self.f1,))     # reopened by hand: same value, allowed
        out = self.post(f'/runs/{rid}/reply', {'message_id': 'ms1', 'sender': 'captain@example.com', 'received_at': '2026-10-06 10:00:00',
                        'updates': [{'finding_id': self.f3, 'remark': 'Wire renewed, completed.', 'close': True},
                                    {'finding_id': self.f1, 'remark': 'Rectified.', 'close': True}]}).get_json()
        self.assertEqual((2, 1), (out['updated'], out['closed']))
        with self.app_context():
            f3 = query('SELECT remark,status FROM cs_findings WHERE id=?', (self.f3,), one=True)
            self.assertEqual(('[10/6]Wire renewed, completed.', 'Closed'), (f3['remark'], f3['status']))   # remark appended, hand status kept
            execute("UPDATE cs_findings SET status='Open' WHERE id=?", (self.f3,))     # reopened by hand after mail
            self.assertIn('수동 상태 변경', query("SELECT note FROM cs_mail_replies WHERE message_id='ms1'", one=True)['note'])
            self.assertEqual('MANUAL', query('SELECT last_status FROM cs_mail_item_state WHERE run_id=? AND finding_id=?', (rid, self.f3), one=True)['last_status'])
        # Hand-reopened f3: status untouched again even if vessel says done.
        self.post(f'/runs/{rid}/reply', {'message_id': 'ms2', 'sender': 'captain@example.com', 'updates': [{'finding_id': self.f3, 'remark': 'Done again.', 'close': True}]})
        with self.app_context():
            self.assertEqual('Open', query('SELECT status FROM cs_findings WHERE id=?', (self.f3,), one=True)['status'])

    def test_runner_pre_send_gate_aborts_when_toggle_turned_off(self):
        import importlib.util
        from pathlib import Path
        source = Path.home() / '.openclaw/workspace/automation/cs-mail/cs_mail_runner.py'
        spec = importlib.util.spec_from_file_location('cs_runner_gate', source)
        runner = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(runner)
        self.client.post(f'/api/cs-mail/settings/{self.vid}/enabled', json={'enabled': True})

        def http(method, path, body=None, **kwargs):
            response = self.client.open(path, method=method, headers=self.headers, json=body)
            return response.status_code, response.get_json()
        clicked = []

        def send(subject, body, to, cc, attachment, before_send=None):
            with self.app_context():   # toggle OFF while waiting for the Outlook lock
                execute('UPDATE cs_mail_settings SET enabled=0 WHERE vessel_id=?', (self.vid,))
            before_send()
            clicked.append(subject)
        with patch.object(runner, 'STATE', Path(self.temp.name)), patch.object(runner.dm, 'http', side_effect=http), \
                patch.object(runner.dm, 'outlook_send', side_effect=send):
            result = runner.send_weekly(SimpleNamespace(dry=False))
        self.assertFalse(result['ok'])
        self.assertEqual([], clicked)
        with self.app_context():
            self.assertEqual('failed', query('SELECT state FROM cs_mail_runs', one=True)['state'])

    # ── reminders ──
    def test_reminder_thursday_once_same_week_only(self):
        rid = self.sent_run('2026-10-05 07:31:00')                     # Monday W41
        for when in (datetime(2026, 10, 7, 8, 30), datetime(2026, 10, 8, 8, 29), datetime(2026, 10, 9, 8, 30), datetime(2026, 10, 15, 8, 30)):
            with patch.object(R, 'now_kst', return_value=when):
                self.assertEqual([], self.get('/reminders').get_json()['runs'], when)
                self.assertEqual(409, self.post(f'/runs/{rid}/reminder', {'kind': 'claim', 'n': 1}).status_code, when)
        with patch.object(R, 'now_kst', return_value=datetime(2026, 10, 8, 8, 30)):     # Thursday
            due = self.get('/reminders').get_json()['runs']
            self.assertEqual([rid], [r['id'] for r in due])
            self.assertEqual(2, len(due[0]['attachments']))
            claim = self.post(f'/runs/{rid}/reminder', {'kind': 'claim', 'n': 1})
            self.assertEqual(200, claim.status_code, claim.get_json())
            token = claim.get_json()['token']
            self.assertEqual(200, self.post(f'/runs/{rid}/reminder', {'kind': 'check', 'n': 1, 'token': token}).status_code)
            self.assertEqual(200, self.post(f'/runs/{rid}/reminder', {'kind': 'result', 'n': 1, 'token': token, 'state': 'sent', 'sent_at': '2026-10-08 08:30:00'}).status_code)
            self.assertEqual(409, self.post(f'/runs/{rid}/reminder', {'kind': 'claim', 'n': 2}).status_code)
            self.assertEqual(409, self.post(f'/runs/{rid}/reminder', {'kind': 'claim', 'n': 1}).status_code)
        with self.app_context():
            self.assertEqual(1, query('SELECT COUNT(*) n FROM cs_mail_reminder_attempts', one=True)['n'])
            self.assertEqual(0, query('SELECT COUNT(*) n FROM mail_reminder_attempts', one=True)['n'])
        self.assertEqual(1, R.max_reminders('cs'))

    def test_reminder_blocked_by_substantive_reply_or_hold_not_by_receipt(self):
        rid = self.sent_run()
        thursday = patch.object(R, 'now_kst', return_value=datetime(2026, 10, 8, 9, 0))
        with thursday:
            # receipt-only replies are never posted by the runner → still due
            self.assertEqual(1, len(self.get('/reminders').get_json()['runs']))
            self.post(f'/runs/{rid}/reply', {'message_id': 'h#hold', 'sender': 'x@example.com', 'updates': [], 'note': '회신 확인 필요(보류): bad rows'})
            self.assertEqual([], self.get('/reminders').get_json()['runs'])
            with self.app_context():
                execute('DELETE FROM cs_mail_replies')
            self.post(f'/runs/{rid}/reply', {'message_id': 'a1', 'sender': 'captain@example.com', 'updates': [{'finding_id': self.f3, 'remark': 'Wire ordered', 'close': False}]})
            self.assertEqual([], self.get('/reminders').get_json()['runs'])

    def test_reminder_not_sent_when_all_sent_findings_closed(self):
        rid = self.sent_run()
        with self.app_context():
            execute("UPDATE cs_findings SET status='Closed'")
        with patch.object(R, 'now_kst', return_value=datetime(2026, 10, 8, 8, 30)):
            self.assertEqual(409, self.post(f'/runs/{rid}/reminder', {'kind': 'claim', 'n': 1}).status_code)

    def test_admin_page_and_template_english(self):
        self.assertEqual(200, self.client.get('/cs-mail').status_code)
        self.assertEqual(401, A.app.test_client().get('/api/cs-mail/status').status_code)
        for body in ('한글', '{unknown}'):
            self.assertEqual(400, self.client.put('/api/cs-mail/template', json={'subject_tpl': 'CS {vessel}', 'body_tpl': body}).status_code)

    # ── runner round trip (no Outlook, no LLM, no network) ──
    def test_runner_send_reply_roundtrip_isolated(self):
        import importlib.util
        from pathlib import Path
        import openpyxl
        source = Path.home() / '.openclaw/workspace/automation/cs-mail/cs_mail_runner.py'
        spec = importlib.util.spec_from_file_location('cs_runner_e2e', source)
        runner = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(runner)
        self.client.post(f'/api/cs-mail/settings/{self.vid}/enabled', json={'enabled': True})

        def http(method, path, body=None, **kwargs):
            response = self.client.open(path, method=method, headers=self.headers, json=body)
            return response.status_code, response.get_json()
        sent = {}

        def capture(subject, body, to, cc, attachment, before_send=None):
            before_send()
            self.assertTrue(subject.startswith('[TRMT-COND '))
            self.assertEqual(['captain@example.com'], to)
            self.assertEqual(2, len(attachment))
            sent['paths'] = attachment
        with patch.object(runner, 'STATE', Path(self.temp.name)), patch.object(runner.dm, 'http', side_effect=http), \
                patch.object(runner.dm, 'outlook_send', side_effect=capture) as send:
            result = runner.send_weekly(SimpleNamespace(dry=False))
            self.assertTrue(result['ok'], result)
            self.assertTrue(runner.send_weekly(SimpleNamespace(dry=False))['ok'])
            self.assertEqual(1, send.call_count)
            reply_path = os.path.join(self.temp.name, 'CS_BOX STAR_2026Q2.xlsx')
            book = openpyxl.load_workbook(next(p for p in sent['paths'] if p.endswith('2026Q2.xlsx')))
            ws = book.active
            ws.cell(5, 5, 'Brake lining renewed on 5 Oct and tested satisfactory.')
            ws.cell(5, 7, 'Closed')
            ws.cell(7, 5, 'Davit wire ordered, will be renewed at Busan.')
            book.save(reply_path)
            run = self.get('/runs/pending').get_json()['runs'][0]
            record = dict(sender='captain@example.com', received_at=run['sent_at'], subject='RE: ' + run['subject'],
                          body='Dear Sir,\nPlease find attached.\n\nFrom: ysson', message_id='o1', attachments=[reply_path])
            verdicts = {self.f1: 'completed', self.f3: 'planned'}
            with patch.object(runner.dm, 'outlook_scan', return_value=([record], '')), \
                    patch.object(runner, '_decision_split', side_effect=lambda r: (dict(r), {})), \
                    patch.object(runner.dm, 'llm', side_effect=lambda prompt, **k: {'items': [{'finding_id': f, 'kind': verdicts[f]} for f in verdicts if f'"finding_id": {f}' in prompt or f'"{f}"' in prompt]}):
                reply = runner.poll(SimpleNamespace(dry=False))
                self.assertTrue(reply['ok'], reply)
                self.assertEqual((2, 1), (reply['replies'][0]['updated'], reply['replies'][0]['closed']))
                self.assertEqual([], runner.poll(SimpleNamespace(dry=False))['replies'])
        with self.app_context():
            self.assertEqual('Closed', query('SELECT status FROM cs_findings WHERE id=?', (self.f1,), one=True)['status'])
            self.assertEqual('Open', query('SELECT status FROM cs_findings WHERE id=?', (self.f3,), one=True)['status'])
            self.assertTrue(query('SELECT remark FROM cs_findings WHERE id=?', (self.f3,), one=True)['remark'].endswith('Davit wire ordered, will be renewed at Busan.'))
            self.assertEqual(1, query('SELECT COUNT(*) n FROM cs_mail_runs', one=True)['n'])


if __name__ == '__main__':
    unittest.main()
