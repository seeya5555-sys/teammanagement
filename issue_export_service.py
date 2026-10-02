"""Daily 업무관리 Excel 워크북 생성 — request 비의존 서비스.

`routes_core.api_issue_export`(화면 엑셀/영문 엑셀 추출)와
`routes_daily_mail`(맥 러너용 선박별 영문 xlsx)이 **같은 빌더**를 쓴다.
두 경로의 시트 구조가 어긋나면 회신 xlsx 파싱이 깨지므로 한 곳에 둔다.

영문(EN) 모드에서만 끝에 회신용 3열을 덧붙인다(기존 9열 위치·의미는 불변 —
기존 소비자가 A~I 열을 위치로 읽어도 깨지지 않게 **뒤에만** 붙인다):
  · `Issue ID`              — TRMT issues.id. 회신 xlsx 를 이슈에 되돌려 붙이는 키.
  · `Update (reply here)`   — 선박이 진행상황을 적는 칸(빈 칸).
  · `Status (Open/Closed)`  — 선박이 Closed 로 바꿔 회신하는 칸(현재 상태로 채움).
"""
from datetime import datetime
from io import BytesIO

#: 회신 파싱이 헤더 **이름**으로 찾는 열. 문자열을 바꾸면 Mac 러너 파서도 같이 바꿔야 한다.
REPLY_ID_HEADER = 'Issue ID'
REPLY_UPDATE_HEADER = 'Update (reply here)'
REPLY_STATUS_HEADER = 'Status (Open/Closed)'

VTYPE_ORDER = ['VLCC', 'LR', 'AFRAMAX', 'MR', 'CNTR']


def sheet_safe(name):
    bad = '[]:*?/\\'
    out = ''.join('_' if c in bad else c for c in (name or ''))
    return (out[:31] or 'Sheet')


def _fmt_actions(acts):
    if not acts:
        return ''
    lines = []
    for a in acts:
        d = (a.get('date') or '').strip()
        p = (a.get('progress') or '').strip()
        mark = '★ ' if a.get('important') else ''
        if d and p:   lines.append(f'{mark}[{d}] {p}')
        elif d:       lines.append(f'{mark}[{d}]')
        elif p:       lines.append(f'{mark}{p}')
    return '\n'.join(lines)


def build_issue_workbook(rows, en=False, sub_text=None, exported_by='', now=None):
    """이슈 dict 리스트(actions 는 이미 list) → (xlsx bytes, 파일명, 선박 시트명 리스트).

    rows 는 `vessel_name`, `vessel_type` 를 포함해야 한다. 번역은 호출부 책임.
    openpyxl 미설치면 ImportError 를 그대로 올린다(호출부가 500 메시지로 바꾼다).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    EN = bool(en)
    now = now or datetime.now()

    def _vrank(t):
        t = (t or '').upper()
        return VTYPE_ORDER.index(t) if t in VTYPE_ORDER else len(VTYPE_ORDER)
    ves_map = {}
    for r in rows:
        vn = r.get('vessel_name') or ('Unassigned' if EN else '미배정')
        if vn not in ves_map:
            ves_map[vn] = {'type': r.get('vessel_type') or '', 'rows': []}
        ves_map[vn]['rows'].append(r)
    ves_seq = sorted(ves_map.keys(), key=lambda n: (_vrank(ves_map[n]['type']), n))

    HEADERS = (['No.', 'Issue Date', 'Item', 'Description', 'Action Plan',
                'Priority', 'Status', 'Due Date', 'TSI Comment',
                REPLY_ID_HEADER, REPLY_UPDATE_HEADER, REPLY_STATUS_HEADER]
               if EN else
               ['No.', '발생일', '현안업무', '상세 내용', '진행사항 (조치 이력)',
                '우선순위', '상태', '마감일', 'TSI Comment'])
    COL_WIDTHS = [5, 12, 30, 40, 44, 12, 11, 12, 34] + ([10, 44, 16] if EN else [])
    N_COLS = len(HEADERS)
    PRI_COL, STAT_COL = 6, 7

    F = 'Malgun Gothic'
    title_font   = Font(name=F, size=14, bold=True, color='FFFFFF')
    sub_font     = Font(name=F, size=10, color='ECF0F1', italic=True)
    title_fill   = PatternFill('solid', start_color='1F3A5F')
    sub_fill     = PatternFill('solid', start_color='2C5282')
    col_hdr_font = Font(name=F, size=10, bold=True, color='FFFFFF')
    col_hdr_fill = PatternFill('solid', start_color='34495E')
    reply_hdr_fill = PatternFill('solid', start_color='6D4C0F')   # 회신 칸은 눈에 띄게(웜 브라운)
    reply_fill   = PatternFill('solid', start_color='FFF8E1')
    body_font    = Font(name=F, size=10)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    body_align   = Alignment(horizontal='left',   vertical='top',    wrap_text=True)
    cent_top     = Alignment(horizontal='center', vertical='top',    wrap_text=True)

    thin = Side(style='thin',   color='BDC3C7')
    med  = Side(style='medium', color='34495E')
    border_thin = Border(left=thin, right=thin, top=thin, bottom=thin)

    PRI_FILL = {
        'COC & Flag': PatternFill('solid', start_color='F8CECC'),
        'Urgent':     PatternFill('solid', start_color='FFE6CC'),
        'Next DD':    PatternFill('solid', start_color='FFF2CC'),
        'Normal':     None,
    }
    PRI_FONT = {
        'COC & Flag': Font(name=F, size=10, bold=True, color='B71C1C'),
        'Urgent':     Font(name=F, size=10, bold=True, color='E65100'),
        'Next DD':    Font(name=F, size=10, bold=True, color='6D4C0F'),
        'Normal':     Font(name=F, size=10, color='5D6D7E'),
    }
    STAT_FILL = {
        'Open':       PatternFill('solid', start_color='E1F5FE'),
        'InProgress': PatternFill('solid', start_color='FFF9C4'),
        'Closed':     PatternFill('solid', start_color='E8F5E9'),
    }
    STAT_FONT = {
        'Open':       Font(name=F, size=10, bold=True, color='0277BD'),
        'InProgress': Font(name=F, size=10, bold=True, color='F57F17'),
        'Closed':     Font(name=F, size=10, bold=True, color='2E7D32'),
    }
    STAT_LABEL = ({'Open': 'Open', 'InProgress': 'In Progress', 'Closed': 'Closed'}
                  if EN else
                  {'Open': 'Open', 'InProgress': '진행중', 'Closed': 'Closed'})
    # 회신 Status 칸은 Open/Closed 2값만 안내한다(In Progress 도 선박 입장에선 Open).
    REPLY_STATUS = {'Open': 'Open', 'InProgress': 'Open', 'Closed': 'Closed'}

    wb = Workbook()
    wb.remove(wb.active)
    today_str = now.strftime('%Y-%m-%d')
    me = exported_by or ''
    sub_text = sub_text or ('All items' if EN else '전체 항목')

    if not ves_seq:
        ws = wb.create_sheet('No Data' if EN else '데이터 없음')
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N_COLS)
        c = ws.cell(row=1, column=1, value=('Daily Work Log — No Data' if EN else 'Daily 업무관리 — 데이터 없음'))
        c.font = title_font; c.fill = title_fill; c.alignment = center_align
        ws.cell(row=3, column=1, value=('No issues match the filter.' if EN else '필터 조건에 해당하는 이슈가 없습니다.')).font = Font(name=F, size=11, italic=True)
        for idx, w in enumerate(COL_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w
    else:
        for vn in ves_seq:
            info = ves_map[vn]
            ws = wb.create_sheet(sheet_safe(vn))
            for idx, w in enumerate(COL_WIDTHS, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = w

            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N_COLS)
            vt = info['type']
            c1 = ws.cell(row=1, column=1, value=(f'{vn}   |   {vt}' if vt else vn))
            c1.font = title_font; c1.fill = title_fill
            c1.alignment = Alignment(horizontal='left', vertical='center', indent=1)
            ws.row_dimensions[1].height = 30

            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=N_COLS)
            cnt = len(info['rows'])
            if EN:
                sub_msg = f'Exported: {today_str}    │    Total {cnt}    │    {sub_text}'
                if me: sub_msg += f'    │    By: {me}'
            else:
                sub_msg = f'추출일: {today_str}    │    총 {cnt}건    │    {sub_text}'
                if me: sub_msg += f'    │    출력: {me}'
            c2 = ws.cell(row=2, column=1, value=sub_msg)
            c2.font = sub_font; c2.fill = sub_fill
            c2.alignment = Alignment(horizontal='left', vertical='center', indent=1)
            ws.row_dimensions[2].height = 20
            ws.row_dimensions[3].height = 6

            HDR_ROW = 4
            for col_idx, h in enumerate(HEADERS, start=1):
                c = ws.cell(row=HDR_ROW, column=col_idx, value=h)
                c.font = col_hdr_font
                c.fill = reply_hdr_fill if col_idx > 9 else col_hdr_fill
                c.alignment = center_align
                c.border = Border(left=thin, right=thin, top=med, bottom=med)
            ws.row_dimensions[HDR_ROW].height = 26

            cur_row = HDR_ROW + 1
            for no, r in enumerate(sorted(info['rows'],
                                          key=lambda x: ((x.get('issue_date') or ''), x.get('id') or 0)), start=1):
                vals = [
                    no,
                    r.get('issue_date') or '',
                    r.get('item_topic') or '',
                    r.get('description') or '',
                    _fmt_actions(r.get('actions')),
                    r.get('priority') or '',
                    STAT_LABEL.get(r.get('status'), r.get('status') or ''),
                    r.get('due_date') or '',
                    '',                                   # TSI Comment — 수기 기입용 빈 칸
                ]
                if EN:
                    vals += [r.get('id'), '', REPLY_STATUS.get(r.get('status'), 'Open')]
                for col_idx, v in enumerate(vals, start=1):
                    c = ws.cell(row=cur_row, column=col_idx, value=v)
                    c.font = body_font
                    c.border = border_thin
                    if col_idx in (1, 2, 8, 10, 12):
                        c.alignment = cent_top
                    elif col_idx in (PRI_COL, STAT_COL):
                        c.alignment = center_align
                    else:
                        c.alignment = body_align
                    if col_idx in (11, 12):
                        c.fill = reply_fill
                pri = r.get('priority')
                if PRI_FILL.get(pri): ws.cell(row=cur_row, column=PRI_COL).fill = PRI_FILL[pri]
                if pri in PRI_FONT:   ws.cell(row=cur_row, column=PRI_COL).font = PRI_FONT[pri]
                st = r.get('status')
                if STAT_FILL.get(st): ws.cell(row=cur_row, column=STAT_COL).fill = STAT_FILL[st]
                if st in STAT_FONT:   ws.cell(row=cur_row, column=STAT_COL).font = STAT_FONT[st]
                cur_row += 1

            last_row = cur_row - 1
            if last_row > HDR_ROW:
                ws.auto_filter.ref = f'A{HDR_ROW}:{get_column_letter(N_COLS)}{last_row}'
            if EN and last_row > HDR_ROW:
                from openpyxl.worksheet.datavalidation import DataValidation
                dv = DataValidation(type='list', formula1='"Open,Closed"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f'L{HDR_ROW + 1}:L{last_row}')
            ws.freeze_panes = f'A{HDR_ROW + 1}'
            ws.print_options.horizontalCentered = True
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_title_rows = f'{HDR_ROW}:{HDR_ROW}'

    today = now.strftime('%Y%m%d')
    suffix = '_EN' if EN else ''
    if len(ves_seq) == 1:
        fname = f'TRMT_Daily_{sheet_safe(ves_seq[0])}_{today}{suffix}.xlsx'
    else:
        fname = f'TRMT_Daily_{today}{suffix}.xlsx'

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue(), fname, ves_seq
