"""Class mail HTTP adapters; admin UI and API-key-only Mac runner."""
import base64
import hashlib
import io
import re

from flask import Blueprint, jsonify, render_template, request, session
from helpers_shared import _class_workbook, admin_required, api_key_required
import class_mail_service as svc
from mail_common import error_response, json_body

bp = Blueprint('routes_class_mail', __name__)


payload = json_body


def call(fn, *args, success=200):
    try:
        return jsonify(fn(*args)), success
    except svc.ClassMailError as error:
        return error_response(error)


@bp.route('/class-mail')
@admin_required
def class_mail_page():
    return render_template('class_mail.html')


@bp.route('/api/class-mail/status')
@admin_required
def api_class_mail_status():
    return jsonify(svc.status())


@bp.route('/api/class-mail/settings/<int:vid>/enabled', methods=['POST'])
@admin_required
def api_class_mail_enabled(vid):
    data = payload()
    if not isinstance(data.get('enabled'), bool):
        return jsonify({'error': 'enabled must be boolean'}), 400
    return call(svc.set_enabled, vid, data['enabled'], session.get('username') or 'admin')


@bp.route('/api/class-mail/runs/<int:rid>/recover-sent', methods=['POST'])
@admin_required
def api_class_mail_recover_sent(rid):
    return call(svc.recover_sent, rid, payload(), session.get('username') or 'admin')


@bp.route('/api/class-mail/template', methods=['PUT'])
@admin_required
def api_class_mail_template():
    return call(svc.save_template, payload(), session.get('username') or 'admin')


@bp.route('/api/ext/class-mail/config')
@api_key_required
def api_ext_class_mail_config():
    return jsonify({'today': svc.today_kst().isoformat(), 'vessels': svc.settings()} if request.args.get('preview') == '1' else svc.config())


@bp.route('/api/ext/class-mail/vessels/<int:vid>/prepare')
@api_key_required
def api_ext_class_mail_prepare(vid):
    """Workbook and exact English row metadata, with source fingerprint.

    Same English-only builder as Class Status's Excel button. No operational write.
    """
    from openpyxl import load_workbook
    vessel = next((v for v in svc.settings() if v['vessel_id'] == vid), None)
    if not vessel or not vessel['items']:
        return jsonify({'error': 'no class findings in roster'}), 404
    items = vessel['items']
    headers = ['Category', 'No', 'Issued', 'Description', 'Due', 'Action Taken', 'Urgent']
    rows = [[svc.CATEGORIES[i['category']], i['no'], i['issued_date'] or '', i['description'] or '',
             i['due_date'] or '', i['action_taken'] or '', i['importance'] or ''] for i in items]
    book = _class_workbook(vessel['name'] + ' Class Status', 'Generated ' + svc.today_kst().isoformat(),
                           headers, rows, {4, 6}, [24, 5, 13, 60, 13, 40, 8])
    if isinstance(book, tuple):
        return book
    raw = book.getvalue()
    ws = load_workbook(io.BytesIO(raw), data_only=True).active
    reply_rows = []
    for item, row in zip(items, ws.iter_rows(min_row=5, values_only=True)):
        reply_rows.append({'item_id': item['id'], 'category': str(row[0] or ''), 'no': str(row[1] or ''),
                           'issued': str(row[2] or ''), 'description': str(row[3] or ''),
                           'due': str(row[4] or ''), 'action': str(row[5] or '')})
    return jsonify({'vessel_id': vid, 'fingerprint': vessel['fingerprint'], 'iso_week': vessel['iso_week'],
                    'filename': re.sub(r'[^A-Za-z0-9_-]+', '_', vessel['name']).strip('_')[:70] + '_ClassStatus_' + vessel['iso_week'] + '.xlsx',
                    'excel_sha256': hashlib.sha256(raw).hexdigest(),
                    'xlsx_base64': base64.b64encode(raw).decode(), 'reply_rows': reply_rows})


@bp.route('/api/ext/class-mail/runs', methods=['POST'])
@api_key_required
def api_ext_class_mail_claim():
    return call(svc.claim, payload(), success=201)


@bp.route('/api/ext/class-mail/canary', methods=['POST'])
@api_key_required
def api_ext_class_mail_canary():
    data = payload()
    if data.get('confirmed_send') is not True:
        return jsonify({'error': 'explicit confirmed_send required'}), 400
    return call(svc.claim, data, True, success=201)


@bp.route('/api/ext/class-mail/runs/<int:rid>/state', methods=['POST'])
@api_key_required
def api_ext_class_mail_state(rid):
    return call(svc.state, rid, payload())


@bp.route('/api/ext/class-mail/runs/pending')
@api_key_required
def api_ext_class_mail_pending():
    return jsonify({'runs': svc.pending()})


@bp.route('/api/ext/class-mail/runs/<int:rid>/reply', methods=['POST'])
@api_key_required
def api_ext_class_mail_reply(rid):
    return call(svc.apply_reply, rid, payload())


@bp.route('/api/ext/class-mail/runs/<int:rid>/check')
@api_key_required
def api_ext_class_mail_check(rid):
    return call(svc.check_send, rid)


@bp.route('/api/ext/class-mail/reminders')
@api_key_required
def api_ext_class_mail_reminders():
    import mail_reminder_service as reminders
    return jsonify(reminders.pending('class'))


@bp.route('/api/ext/class-mail/runs/<int:rid>/reminder', methods=['POST'])
@api_key_required
def api_ext_class_mail_reminder(rid):
    import mail_reminder_service as reminders
    try:
        return jsonify(reminders.action('class', rid, json_body()))
    except reminders.ReminderError as error:
        return error_response(error)
